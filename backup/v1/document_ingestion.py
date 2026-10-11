"""Extract and lightly preprocess text from supported document formats."""

from io import BytesIO
from pathlib import Path
import re

from docx import Document
from pypdf import PdfReader


TEXT_EXTENSIONS = ("pdf", "docx", "txt")
IMAGE_EXTENSIONS = ("png", "jpg", "jpeg", "bmp", "tif", "tiff", "webp", "gif")
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS + IMAGE_EXTENSIONS


def _preprocess_text(text: str) -> str:
    """Normalize whitespace without changing the words or paragraph boundaries."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n[ \t]+", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extract_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(BytesIO(file_bytes))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(file_bytes: bytes) -> str:
    document = Document(BytesIO(file_bytes))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def extract_text(file_name: str, file_bytes: bytes) -> str:
    """Extract text from a PDF, DOCX, or TXT file and normalize its whitespace.

    Raises:
        ValueError: If the file extension is not supported.
    """
    extension = Path(file_name).suffix.lower().lstrip(".")

    if extension == "pdf":
        text = _extract_pdf(file_bytes)
    elif extension == "docx":
        text = _extract_docx(file_bytes)
    elif extension == "txt":
        text = file_bytes.decode("utf-8-sig", errors="replace")
    else:
        supported = ", ".join(TEXT_EXTENSIONS)
        raise ValueError(
            f"Direct text extraction supports: {supported}. Use the OCR module for images."
        )

    return _preprocess_text(text)